"""Safe document parsing. Extracts text only - nothing inside documents is ever executed."""
import io
from dataclasses import dataclass
from pathlib import Path

DOCUMENT_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".txt": "text/plain",
    ".md": "text/markdown",
}
IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif"}


class UnsupportedFile(ValueError):
    pass


@dataclass
class ValidatedFile:
    extension: str
    content_type: str
    kind: str  # document | image


def validate_upload(filename: str, data: bytes, max_bytes: int, allow_images: bool = True, allow_documents: bool = True) -> ValidatedFile:
    """Validates by extension AND file signature (magic bytes) and size."""
    if not data:
        raise UnsupportedFile("The file is empty.")
    if len(data) > max_bytes:
        raise UnsupportedFile(f"File is too large (max {max_bytes // (1024 * 1024)} MB).")
    ext = Path(filename or "").suffix.lower()
    if allow_documents and ext in DOCUMENT_TYPES:
        if ext == ".pdf" and not data.startswith(b"%PDF"):
            raise UnsupportedFile("File content is not a valid PDF.")
        if ext in (".docx", ".xlsx") and not data.startswith(b"PK\x03\x04"):
            raise UnsupportedFile(f"File content is not a valid {ext} document.")
        if ext in (".csv", ".txt", ".md"):
            if b"\x00" in data[:8192]:
                raise UnsupportedFile("Text file appears to be binary.")
        return ValidatedFile(ext, DOCUMENT_TYPES[ext], "document")
    if allow_images and ext in IMAGE_TYPES:
        sig_ok = (
            (ext == ".png" and data.startswith(b"\x89PNG\r\n\x1a\n"))
            or (ext in (".jpg", ".jpeg") and data.startswith(b"\xff\xd8\xff"))
            or (ext == ".gif" and data[:6] in (b"GIF87a", b"GIF89a"))
            or (ext == ".webp" and data[:4] == b"RIFF" and data[8:12] == b"WEBP")
        )
        if not sig_ok:
            raise UnsupportedFile("File content does not match its image type.")
        return ValidatedFile(ext, IMAGE_TYPES[ext], "image")
    allowed = list(DOCUMENT_TYPES if allow_documents else {}) + list(IMAGE_TYPES if allow_images else {})
    raise UnsupportedFile(f"Unsupported file type '{ext or 'unknown'}'. Allowed: {', '.join(allowed)}")


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> str:
    import fitz  # PyMuPDF

    parts = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        if doc.needs_pass:
            raise UnsupportedFile("Password-protected PDFs are not supported.")
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                parts.append(f"[Page {i}]\n\n{text}")
    return "\n\n".join(parts)


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if p.style is not None and p.style.name.lower().startswith("heading"):
            parts.append(f"## {t}")
        else:
            parts.append(t)
    for table in d.tables:
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append("\n".join(rows))
    return "\n\n".join(parts)


def _frame_to_text(df, title: str) -> str:
    df = df.dropna(how="all").dropna(axis=1, how="all")
    lines = [f"[Sheet: {title}]"]
    cols = [str(c) for c in df.columns]
    for _, row in df.iterrows():
        pairs = [f"{c}: {v}" for c, v in zip(cols, row.tolist()) if str(v) not in ("nan", "None", "")]
        if pairs:
            lines.append("; ".join(pairs))
    return "\n".join(lines)


def _xlsx(data: bytes) -> str:
    import pandas as pd

    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, engine="openpyxl")
    return "\n\n".join(_frame_to_text(df, name) for name, df in sheets.items())


def _csv(data: bytes) -> str:
    import pandas as pd

    df = pd.read_csv(io.StringIO(_decode(data)), sep=None, engine="python")
    return _frame_to_text(df, "CSV")


def extract_text(extension: str, data: bytes) -> str:
    ext = extension.lower()
    if ext == ".pdf":
        text = _pdf(data)
    elif ext == ".docx":
        text = _docx(data)
    elif ext == ".xlsx":
        text = _xlsx(data)
    elif ext == ".csv":
        text = _csv(data)
    elif ext in (".txt", ".md"):
        text = _decode(data)
    else:
        raise UnsupportedFile(f"Cannot extract text from {ext}")
    text = text.replace("\x00", "").strip()
    if not text:
        raise UnsupportedFile("No extractable text found (scanned PDFs need OCR before upload).")
    return text
