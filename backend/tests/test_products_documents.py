import io

from sqlalchemy import func, select

from app.models import DocumentChunk
from tests.conftest import SAMPLE_PRODUCT

PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64


def chunk_count(db, source_type, source_id):
    return db.scalar(select(func.count()).select_from(DocumentChunk).where(DocumentChunk.source_type == source_type, DocumentChunk.source_id == source_id))


# ---------------- products ----------------
def test_product_crud_versioning_and_index_sync(client, admin_h, product, db):
    pid = product["id"]
    assert product["version"] == 1
    assert chunk_count(db, "product", pid) > 0
    # Specification order is preserved
    assert list(product["specifications"]) == list(SAMPLE_PRODUCT["specifications"])

    upd = {**SAMPLE_PRODUCT, "warranty": "Now 48 months warranty.", "change_note": "Warranty extended"}
    r = client.put(f"/api/products/{pid}", json=upd, headers=admin_h)
    assert r.status_code == 200 and r.json()["version"] == 2
    versions = client.get(f"/api/products/{pid}/versions", headers=admin_h).json()
    assert [v["version"] for v in versions] == [2, 1]
    assert versions[0]["change_note"] == "Warranty extended"

    search = client.post("/api/knowledge/search", json={"query": "HydroMax 900 warranty months"}, headers=admin_h).json()
    assert any("48 months" in r["content"] for r in search["results"])
    assert not any("30 month" in r["content"] for r in search["results"])

    r = client.post(f"/api/products/{pid}/versions/1/restore", headers=admin_h)
    assert r.json()["version"] == 3 and "30 month" in r.json()["warranty"]

    client.post(f"/api/products/{pid}/archive", headers=admin_h)
    db.expire_all()
    assert chunk_count(db, "product", pid) == 0

    assert client.delete(f"/api/products/{pid}", headers=admin_h).status_code == 204
    assert client.get(f"/api/products/{pid}", headers=admin_h).status_code == 404


def test_product_list_search_filter_pagination(client, admin_h):
    cat = client.post("/api/categories", json={"name": "Pumps"}, headers=admin_h).json()
    for i in range(5):
        client.post("/api/products", json={**SAMPLE_PRODUCT, "name": f"Pump {i}", "model_number": f"P-{i}", "category_id": cat["id"]}, headers=admin_h)
    client.post("/api/products", json={**SAMPLE_PRODUCT, "name": "Other", "model_number": "O-1"}, headers=admin_h)
    r = client.get("/api/products", params={"category_id": cat["id"], "page_size": 2}, headers=admin_h).json()
    assert r["total"] == 5 and len(r["items"]) == 2
    r = client.get("/api/products", params={"q": "O-1"}, headers=admin_h).json()
    assert r["total"] == 1 and r["items"][0]["name"] == "Other"
    cats = client.get("/api/categories", headers=admin_h).json()
    assert cats[0]["product_count"] == 5


def test_product_image_upload_validation(client, admin_h, product):
    pid = product["id"]
    bad = client.post(f"/api/products/{pid}/images", files={"file": ("x.png", b"not really a png", "image/png")}, headers=admin_h)
    assert bad.status_code == 400
    exe = client.post(f"/api/products/{pid}/images", files={"file": ("x.exe", b"MZ....", "application/octet-stream")}, headers=admin_h)
    assert exe.status_code == 400
    ok = client.post(f"/api/products/{pid}/images", files={"file": ("photo.png", PNG, "image/png")}, headers=admin_h)
    assert ok.status_code == 201
    url = ok.json()["url"]
    assert client.get(url).status_code == 200  # signed URL works without auth header
    assert client.get(url.replace("sig=", "sig=0")).status_code == 403


# ---------------- documents ----------------
def _docx_bytes() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Installation Guide", level=1)
    d.add_paragraph("The Zephyr X2 fan must be installed at least 2.3 metres above the floor.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Blade span", "1200 mm"
    t.cell(1, 0).text, t.cell(1, 1).text = "Speed", "380 RPM"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prices"
    ws.append(["Model", "Dealer Price", "Stock"])
    ws.append(["Zephyr X2", 2450, 31])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf_bytes() -> bytes:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Zephyr X2 user manual. Clean the blades every 30 days with a dry cloth.")
    return doc.tobytes()


def test_upload_and_parse_all_formats(client, admin_h, db):
    cases = [
        ("guide.docx", _docx_bytes(), "2.3 metres"),
        ("prices.xlsx", _xlsx_bytes(), "Dealer Price: 2450"),
        ("manual.pdf", _pdf_bytes(), "every 30 days"),
        ("stock.csv", b"model,colour,stock\nZephyr X2,white,31\n", "colour: white"),
        ("notes.txt", "Zephyr X2 ships with a remote control. Ünïcödé ok.".encode(), "remote control"),
        ("faq.md", b"# FAQ\n\nThe Zephyr X2 works with inverters.", "works with inverters"),
    ]
    for name, data, expected in cases:
        r = client.post("/api/documents/upload", files={"file": (name, data)}, data={"category": "user_manuals"}, headers=admin_h)
        assert r.status_code == 201, (name, r.text)
        doc = client.get(f"/api/documents/{r.json()['id']}", headers=admin_h).json()
        assert doc["processing_status"] == "ready", (name, doc["processing_error"])
        assert expected in doc["content"], (name, doc["content"][:300])
        assert doc["chunk_count"] >= 1
        assert doc["download_url"]


def test_upload_rejects_bad_files(client, admin_h):
    assert client.post("/api/documents/upload", files={"file": ("fake.pdf", b"hello", "application/pdf")}, headers=admin_h).status_code == 400
    assert client.post("/api/documents/upload", files={"file": ("x.docx", b"not a zip", "application/octet-stream")}, headers=admin_h).status_code == 400
    assert client.post("/api/documents/upload", files={"file": ("run.sh", b"#!/bin/sh\nrm -rf /", "text/x-sh")}, headers=admin_h).status_code == 400
    assert client.post("/api/documents/upload", files={"file": ("empty.txt", b"", "text/plain")}, headers=admin_h).status_code == 400


def test_upload_size_limit(client, admin_h, monkeypatch):
    from app.core.config import get_settings

    monkeypatch.setattr(get_settings(), "max_upload_mb", 1)
    big = b"a" * (1024 * 1024 + 10)
    assert client.post("/api/documents/upload", files={"file": ("big.txt", big)}, headers=admin_h).status_code == 400


def test_document_versions_update_and_delete_sync(client, admin_h, db):
    r = client.post("/api/documents", json={"title": "Returns policy", "content": "Customers may return goods within 10 days.", "category": "company_policies"}, headers=admin_h)
    doc_id = r.json()["id"]
    r = client.patch(f"/api/documents/{doc_id}", json={"content": "Customers may return goods within 15 days.", "change_note": "Policy update"}, headers=admin_h)
    assert r.json()["current_version"] == 2
    res = client.post("/api/knowledge/search", json={"query": "return goods days"}, headers=admin_h).json()["results"]
    assert any("15 days" in x["content"] for x in res) and not any("10 days" in x["content"] for x in res)
    # New file version
    r = client.post(f"/api/documents/{doc_id}/versions", files={"file": ("policy.txt", b"Customers may return goods within 20 days.")}, data={"change_note": "v3"}, headers=admin_h)
    assert r.json()["current_version"] == 3 and len(r.json()["versions"]) == 3
    restored = client.post(f"/api/documents/{doc_id}/versions/1/restore", headers=admin_h).json()
    assert "10 days" in restored["content"] and restored["current_version"] == 4
    # Archive removes from index; delete removes the document
    client.patch(f"/api/documents/{doc_id}", json={"status": "archived"}, headers=admin_h)
    db.expire_all()
    assert chunk_count(db, "document", doc_id) == 0
    assert client.delete(f"/api/documents/{doc_id}", headers=admin_h).status_code == 204
    assert client.get(f"/api/documents/{doc_id}", headers=admin_h).status_code == 404


def test_audit_log_records_changes(client, admin_h, product):
    logs = client.get("/api/audit-logs", headers=admin_h).json()["items"]
    assert any(l["action"] == "product.create" and l["entity_id"] == str(product["id"]) for l in logs)


def test_faq_crud_and_retrieval(client, admin_h, member_h):
    r = client.post("/api/faqs", json={"question": "Is there a toll-free helpline?", "answer": "Yes, call 1800-123-999 between 9am and 6pm."}, headers=admin_h)
    fid = r.json()["id"]
    res = client.post("/api/knowledge/search", json={"query": "toll-free helpline number"}, headers=member_h).json()["results"]
    assert res[0]["source_type"] == "faq"
    client.put(f"/api/faqs/{fid}", json={"question": "Is there a toll-free helpline?", "answer": "Yes, call 1800-555-000."}, headers=admin_h)
    res = client.post("/api/knowledge/search", json={"query": "toll-free helpline number"}, headers=member_h).json()["results"]
    assert "1800-555-000" in res[0]["content"]
    client.delete(f"/api/faqs/{fid}", headers=admin_h)
    res = client.post("/api/knowledge/search", json={"query": "toll-free helpline number"}, headers=member_h).json()["results"]
    assert not any(r["source_type"] == "faq" for r in res)
